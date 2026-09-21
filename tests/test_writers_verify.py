import io
import sys
import types

import docx
import pymupdf
import pytest
from PIL import Image as PILImage

from pii_redact.model import Doc, Entity, Page, Span, Word
from pii_redact.verify import verify
from pii_redact.writers.docx_writer import write_docx
from pii_redact.writers.image_redact import redact_image
from pii_redact.writers.pdf_to_docx import pdf_to_docx
from pii_redact.writers.pdf_writer import write_pdf

NAME, PHONE = "Sarthak Malvadkar", "+91 20 4505 3237"
FAKE_NAME, FAKE_PHONE = "Rohan Deshpande", "+91 20 9182 7364"


def _cfg(legend):
    return {"surrogates": {"legend": legend}}


def _entities(words):
    """Build Words -> Spans -> Entities for NAME and PHONE from a list of (id, text, loc, block, line)."""
    def span(needle, typ):
        toks = needle.split()
        for i in range(len(words) - len(toks) + 1):
            if [w.text.strip(",") for w in words[i:i + len(toks)]] == toks:
                ids = [w.id for w in words[i:i + len(toks)]]
                return Span(page=1, start=0, end=0, text=needle, type=typ, finder="regex",
                            confidence=1.0, view="raw", word_ids=ids)
        raise AssertionError(f"{needle!r} not found in words")

    p = Entity(id="P1", type="PERSON", canonical=NAME, surrogate=FAKE_NAME,
               name_tokens=dict(zip(NAME.split(), FAKE_NAME.split())))
    p.mentions.append(span(NAME, "PERSON"))
    t = Entity(id="PHONE1", type="PHONE", canonical=PHONE, surrogate=FAKE_PHONE)
    t.mentions.append(span(PHONE, "PHONE"))
    return [p, t]


# ---------- PDF ----------

@pytest.fixture
def pdf_doc(tmp_path):
    src = tmp_path / "src.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Draft Red Herring Prospectus", fontsize=14)
    page.insert_text((72, 100), f"Contact: {NAME}, Tel {PHONE}", fontsize=11)
    page.insert_text((72, 120), "Registered office: Pune, Maharashtra", fontsize=11)
    # a ruled 2x2 table for pdf_to_docx
    for x in (72, 200, 328):
        page.draw_line((x, 160), (x, 200))
    for y in (160, 180, 200):
        page.draw_line((72, y), (328, y))
    for (x, y), t in {(80, 174): "Bank", (208, 174): "Branch", (80, 194): "ICICI", (208, 194): "Kothrud"}.items():
        page.insert_text((x, y), t, fontsize=9)
    pdf.set_metadata({"author": "Sarthak"})
    pdf.save(src)
    pdf.close()

    words = []
    with pymupdf.open(src) as pdf:
        for i, (x0, y0, x1, y1, txt, b, l, _) in enumerate(pdf[0].get_text("words")):
            words.append(Word(id=i, page=1, text=txt, loc={"bbox": (x0, y0, x1, y1)}, block=b, line=l))
    doc = Doc(doc_id="t", path=str(src), fmt="pdf", pages=[Page(page_no=1, words=words)])
    doc.entities = _entities(words)
    return doc, src


@pytest.mark.parametrize("legend,extra_pages", [("none", 0), ("full", 1)])
def test_write_pdf(pdf_doc, tmp_path, legend, extra_pages):
    doc, src = pdf_doc
    out = tmp_path / "out.pdf"
    write_pdf(doc, str(src), str(out), _cfg(legend))

    with pymupdf.open(out) as pdf:
        assert len(pdf) == 1 + extra_pages
        body = pdf[0].get_text()
        assert pdf.metadata.get("author") in (None, "")
    assert FAKE_NAME in body and FAKE_PHONE in body
    assert NAME not in body and PHONE not in body
    assert "Registered office: Pune" in body
    assert verify(doc, str(out), _cfg(legend)) == {"ok": True, "leaks": []}


def test_verify_reports_leak(pdf_doc):
    doc, src = pdf_doc
    res = verify(doc, str(src), _cfg("none"))
    assert not res["ok"]
    assert {(l["page"], l["entity_id"]) for l in res["leaks"]} == {(1, "P1"), (1, "PHONE1")}


def test_pdf_to_docx(pdf_doc, tmp_path):
    doc, src = pdf_doc
    out_pdf, out_docx = tmp_path / "out.pdf", tmp_path / "out.docx"
    write_pdf(doc, str(src), str(out_pdf), _cfg("none"))
    pdf_to_docx(str(out_pdf), str(out_docx))
    d = docx.Document(out_docx)
    text = "\n".join(p.text for p in d.paragraphs)
    assert FAKE_NAME in text and NAME not in text
    assert "Draft Red Herring Prospectus" in text
    assert len(d.tables) == 1
    assert [c.text for c in d.tables[0].rows[1].cells] == ["ICICI", "Kothrud"]
    assert "ICICI" not in text  # table cells are not duplicated as paragraphs


# ---------- DOCX ----------

@pytest.fixture
def docx_doc(tmp_path, monkeypatch):
    src = tmp_path / "src.docx"
    d = docx.Document()
    d.add_paragraph("Draft Red Herring Prospectus")
    p = d.add_paragraph()
    p.add_run("Contact: Sarthak ").bold = True      # run 0 (bold)
    p.add_run(f"Malvadkar, Tel {PHONE}").bold = False  # run 1
    d.core_properties.author = "Sarthak"
    d.save(src)

    def iter_docx_paragraphs(document):
        return list(document.paragraphs)

    monkeypatch.setitem(sys.modules, "pii_redact.parse",
                        types.SimpleNamespace(iter_docx_paragraphs=iter_docx_paragraphs))

    words, wid = [], 0
    for pi, para in enumerate(iter_docx_paragraphs(docx.Document(src))):
        for ri, run in enumerate(para.runs):
            off = 0
            for tok in run.text.split(" "):
                if tok:
                    words.append(Word(id=wid, page=1, text=tok, loc={"para": pi, "run": ri, "offset": off}))
                    wid += 1
                off += len(tok) + 1
    doc = Doc(doc_id="t", path=str(src), fmt="docx", pages=[Page(page_no=1, words=words)])
    doc.entities = _entities(words)
    return doc, src


@pytest.mark.parametrize("legend", ["none", "surrogates", "full"])
def test_write_docx(docx_doc, tmp_path, legend):
    doc, src = docx_doc
    out = tmp_path / "out.docx"
    write_docx(doc, str(src), str(out), _cfg(legend))

    d = docx.Document(out)
    para = d.paragraphs[1]
    assert [r.bold for r in para.runs] == [True, False]
    assert para.runs[0].text == f"Contact: {FAKE_NAME},"
    assert para.runs[1].text == f" Tel {FAKE_PHONE}"
    assert d.core_properties.author == ""
    assert len(d.tables) == (1 if legend != "none" else 0)
    assert verify(doc, str(out), _cfg(legend)) == {"ok": True, "leaks": []}


# ---------- image ----------

def test_redact_image_roundtrip():
    img = PILImage.new("RGB", (200, 80), "gray")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    out = redact_image(buf.getvalue(), [((10, 10, 150, 40), "XXXX1234")])
    with PILImage.open(io.BytesIO(out)) as res:
        assert res.format == "PNG" and res.size == (200, 80)
        assert res.getpixel((12, 12)) == (255, 255, 255)
