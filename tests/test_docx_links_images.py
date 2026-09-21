"""DOCX defects found in review: hyperlink runs/targets bypassed detection; header/footer
images were resolved through the body part. Full pipeline, no LLM."""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

from pii_redact import pipeline
from pii_redact.parse import iter_docx_paragraphs, iter_docx_runs, parse_document

EMAIL = "sarthak.malvadkar@kshinternational.com"
PAN = "NBWPS1951N"


def _add_hyperlink(paragraph, url: str, text: str) -> str:
    """python-docx has no hyperlink API: add the external rel and build <w:hyperlink><w:r><w:t/>."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    r, t = OxmlElement("w:r"), OxmlElement("w:t")
    t.text = text
    r.append(t)
    hl.append(r)
    paragraph._p.append(hl)
    return rid


def _png(text: str) -> bytes:
    img = Image.new("RGB", (520, 90), "white")
    ImageDraw.Draw(img).text((10, 20), text, fill="black", font=ImageFont.load_default(size=40))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _zip_text(path, name: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf8")


def test_hyperlink_text_and_target_are_redacted(tmp_path):
    src = tmp_path / "links.docx"
    d = docx.Document()
    d.add_paragraph("Draft Red Herring Prospectus")
    p = d.add_paragraph("Contact Person: Sarthak Malvadkar, E-mail: ")
    _add_hyperlink(p, f"mailto:{EMAIL}", EMAIL)
    d.save(src)

    # the parser sees the hyperlink run and the mailto: target
    doc = parse_document(str(src), {})
    words = doc.pages[0].words
    assert [w.text for w in words if w.text == EMAIL] == [EMAIL, EMAIL]   # visible run + rel target
    rel_word = next(w for w in words if "rel" in w.loc)
    assert rel_word.loc["rel"].startswith("rId")
    run_word = next(w for w in words if w.text == EMAIL and "run" in w.loc)
    para = iter_docx_paragraphs(docx.Document(src))[run_word.loc["para"]]
    assert iter_docx_runs(para)[run_word.loc["run"]].text == EMAIL
    assert len(para.runs) < len(iter_docx_runs(para))   # Paragraph.runs skips the hyperlink run

    summary = pipeline.run(str(src), str(tmp_path), use_llm=False, use_review=False)
    assert summary["verify"]["ok"], summary["verify"]["leaks"]
    out = summary["output"]
    body, rels = _zip_text(out, "word/document.xml"), _zip_text(out, "word/_rels/document.xml.rels")
    body = body.split("Redaction legend")[0]   # legend=full appends the real values in a table by design
    assert EMAIL not in body and "kshinternational" not in body
    assert EMAIL not in rels and "kshinternational" not in rels
    assert "<w:hyperlink" in body and 'Target="mailto:' in rels and "@example.com" in rels
    assert "Draft Red Herring Prospectus" in body


def test_header_image_is_redacted_through_header_part(tmp_path):
    src = tmp_path / "header_image.docx"
    d = docx.Document()
    d.add_paragraph("Draft Red Herring Prospectus. PAN of the promoter is printed in the header.")
    png = _png(f"PAN {PAN}")
    d.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(png))
    d.save(src)

    doc = parse_document(str(src), {})
    images = [im for p in doc.pages for im in p.images]
    assert len(images) == 1 and images[0].ref.startswith("rId") and "@" in images[0].ref

    summary = pipeline.run(str(src), str(tmp_path), use_llm=False, use_review=False)
    assert summary["verify"]["ok"], summary["verify"]["leaks"]
    spans = list(csv.DictReader(open(Path(summary["output"]).parent / "spans.csv")))
    assert any(r["type"] == "PAN" and r["text"] == PAN for r in spans), spans

    with zipfile.ZipFile(src) as z:
        (media,) = [n for n in z.namelist() if n.startswith("word/media/")]
        before = z.read(media)
    with zipfile.ZipFile(summary["output"]) as z:
        after = z.read(media)
        header_rels = [n for n in z.namelist() if "header" in n and n.endswith(".rels")]
        assert any(media.split("/")[-1] in z.read(n).decode() for n in header_rels)
    assert before == png and after != before

