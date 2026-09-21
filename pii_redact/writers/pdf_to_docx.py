"""Approximate PDF -> DOCX conversion of the redacted PDF (assignment deliverable)."""
from __future__ import annotations

import docx
import pymupdf


def pdf_to_docx(redacted_pdf_path: str, out_path: str) -> None:
    document = docx.Document()
    with pymupdf.open(redacted_pdf_path) as pdf:
        for pno, page in enumerate(pdf):
            if page.get_text().lstrip().startswith("Redaction legend"):
                continue                                    # the real->fake legend must not travel with the DOCX
            if pno:
                document.add_page_break()
            items = []  # (y0, x0, kind, payload)
            table_rects = []
            for tab in page.find_tables().tables:
                rect = pymupdf.Rect(tab.bbox)
                table_rects.append(rect)
                items.append((rect.y0, rect.x0, "table", tab.extract()))
            for x0, y0, x1, y1, text, _, btype in page.get_text("blocks"):
                r = pymupdf.Rect(x0, y0, x1, y1)
                if btype == 0 and text.strip() and not any(r.intersects(t) for t in table_rects):
                    items.append((y0, x0, "para", " ".join(text.split())))
            for _, _, kind, payload in sorted(items, key=lambda t: (t[0], t[1])):
                if kind == "para":
                    document.add_paragraph(payload)
                elif payload:
                    table = document.add_table(rows=len(payload), cols=max(len(r) for r in payload))
                    table.style = "Table Grid"
                    for r, row in enumerate(payload):
                        for c, val in enumerate(row):
                            table.cell(r, c).text = " ".join(str(val or "").split())
    document.save(out_path)
