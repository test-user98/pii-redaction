"""DOCX -> DOCX: replace PII inside runs in place, keeping run formatting."""
from __future__ import annotations

import re

import docx

from pii_redact.model import Doc, Word
from pii_redact.writers._common import LEGEND_HEADING, keep_punctuation, legend_rows, mentions_by_page
from pii_redact.writers.image_redact import redact_image

TO_END = 10**9  # slice end meaning "through the end of the run"


def write_docx(doc: Doc, src_path: str, out_path: str, cfg: dict) -> None:
    from pii_redact.parse import iter_docx_paragraphs, iter_docx_runs

    document = docx.Document(src_path)
    paragraphs = iter_docx_paragraphs(document)
    words = {(p.page_no, w.id): w for p in doc.pages for w in p.words}
    images = {img.ref: img for p in doc.pages for img in p.images if img.ref}

    edits = []  # (para, run, start, end, replacement)
    image_boxes: dict[str, list] = {}  # image ref -> [(bbox_px, surrogate)]
    rel_edits: dict[tuple[int, str], tuple[str, str]] = {}  # (para, rId) -> (span text, surrogate)
    for page_no, mentions in mentions_by_page(doc).items():
        for span, surrogate in mentions:
            span_words = [words[(page_no, i)] for i in span.word_ids if (page_no, i) in words]
            ocr_words = [w for w in span_words if w.source == "ocr"]
            if ocr_words:
                bb = [w.loc["bbox"] for w in ocr_words]
                box = (min(b[0] for b in bb), min(b[1] for b in bb), max(b[2] for b in bb), max(b[3] for b in bb))
                image_boxes.setdefault(ocr_words[0].loc["image_ref"], []).append((box, surrogate))
            for w in span_words:
                if "rel" in w.loc:
                    rel_edits[(w.loc["para"], w.loc["rel"])] = (span.text, surrogate)
            text_words = [w for w in span_words if w.source != "ocr" and "rel" not in w.loc]
            edits.extend(_run_edits(text_words, keep_punctuation(span, text_words, surrogate)))

    for para, run, start, end, repl in sorted(edits, reverse=True):
        r = iter_docx_runs(paragraphs[para])[run]
        r.text = r.text[:start] + repl + r.text[end:]

    for (para, rid), (real, surrogate) in rel_edits.items():
        # Hyperlink target: swap the PII inside the URL, keeping the scheme (mailto:, tel:, https://...).
        rel = paragraphs[para].part.rels[rid]
        target = rel.target_ref
        m = re.match(r"^(mailto|tel):", target, re.I)
        rel._target = target.replace(real, surrogate) if real in target else (m.group(0) if m else "") + surrogate

    by_part: dict[int, tuple] = {}  # one redaction per image part, even when several paragraphs embed it
    for ref, boxes in image_boxes.items():
        rid, para = ref.split("@")  # "<rId>@<para>": the image belongs to that paragraph's part (body/header/footer)
        part = paragraphs[int(para)].part.related_parts[rid]
        by_part.setdefault(id(part), (part, images[ref].png, []))[2].extend(boxes)
    for part, png, boxes in by_part.values():
        part._blob = redact_image(png, boxes)
        part._content_type = "image/png"

    cp = document.core_properties
    cp.author = cp.last_modified_by = cp.comments = cp.title = ""

    rows = legend_rows(doc, cfg)
    if rows:
        document.add_page_break()
        document.add_paragraph().add_run(LEGEND_HEADING).bold = True
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                table.cell(r, c).text = val
    document.save(out_path)


def _run_edits(span_words: list[Word], surrogate: str) -> list[tuple]:
    """Group a span's words by (para, run): first run gets the surrogate, the rest lose the covered slice.
    A span crossing runs covers the first run to its end and later runs from their start."""
    groups: dict[tuple[int, int], list[Word]] = {}
    for w in span_words:
        groups.setdefault((w.loc["para"], w.loc["run"]), []).append(w)
    edits, last = [], len(groups) - 1
    for i, ((para, run), ws) in enumerate(sorted(groups.items())):
        start = min(w.loc["offset"] for w in ws) if i == 0 else 0
        end = max(w.loc["offset"] + len(w.text) for w in ws) if i == last else TO_END
        edits.append((para, run, start, end, surrogate if i == 0 else ""))
    return edits
